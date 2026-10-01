import io
import re
from datetime import datetime

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from src.api import app as api
from src.services.analytics import get_incidence, get_status, get_year_comparison
from src.services.export.excel import build_excel
from src.services.export.figure import build_figure
from src.services.export.model import (
    ChartExport,
    Kpi,
    KpiGroup,
    format_de,
    format_percent,
    format_value,
    trend_export,
    year_export,
)
from src.services.export import pdf as pdf_module
from src.services.export.pdf import build_pdf

NOW = datetime(2026, 10, 1, 12, 0)


def _chart(**overrides) -> ChartExport:
    values = {
        "title": "Titel – Zusatz",
        "subtitle": "Untertitel",
        "x_title": "Kalenderwoche",
        "y_title": "pro 100.000",
        "categories": ["2026-W01", "2026-W02", "2026-W03", "2026-W04"],
        "months": ["JANUAR", "JANUAR", "JANUAR", "JANUAR"],
        "series": {"A": [100.0, None, 120.0, 130.0], "B": [100.5, 105.25, None, 125.0]},
        "kpi_groups": [KpiGroup("Aktuelle Woche", [Kpi("Aktuell", "130", "KW04 · je 100.000 Einw."), Kpi("Vorwoche", "120")])],
        "info": [("Lizenz", "CC BY 4.0")],
        "source_note": "Quelle: RKI",
        "filename": "test",
    }
    values.update(overrides)
    return ChartExport(**values)


def _page_count(pdf: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page(?![A-Za-z])", pdf))


def test_german_number_format() -> None:
    assert format_de(1234.5, 1) == "1.234,5"
    assert format_de(1003) == "1.003"
    assert format_de(None) == "—"
    assert format_percent(4.5) == "+4,5 %"
    assert format_percent(-8.4) == "-8,4 %"
    assert format_percent(None) == "—"
    assert format_value(1216.0) == "1.216"
    assert format_value(915.8) == "915,8"


def test_excel_has_month_and_week_columns_data_blanks_chart_and_info_sheet() -> None:
    workbook = load_workbook(io.BytesIO(build_excel(_chart())))

    assert workbook.sheetnames == ["Daten", "Info", "Kennzahlen"]
    sheet = workbook["Daten"]
    assert [c.value for c in sheet[4]][:4] == ["Monat", "Kalenderwoche", "A", "B"]
    assert [sheet.cell(row=r, column=1).value for r in range(5, 9)] == ["JANUAR"] * 4
    assert [sheet.cell(row=r, column=2).value for r in range(5, 9)] == ["2026-W01", "2026-W02", "2026-W03", "2026-W04"]
    assert [sheet.cell(row=r, column=3).value for r in range(5, 9)] == [100.0, None, 120.0, 130.0]
    assert [sheet.cell(row=r, column=4).value for r in range(5, 9)] == [100.5, 105.25, None, 125.0]
    assert len(sheet._charts) == 1

    info = {row[0].value: row[1].value for row in workbook["Info"].iter_rows(min_row=1) if row[0].value}
    assert info["Lizenz"] == "CC BY 4.0"
    figures = [[c.value for c in row] for row in workbook["Kennzahlen"].iter_rows(min_row=4, max_row=6)]
    assert figures[0] == ["Bereich", "Kennzahl", "Wert", "Bezug"]
    assert figures[1] == ["Aktuelle Woche", "Aktuell", "130", "KW04 · je 100.000 Einw."]


def test_excel_summary_sheet_holds_real_numbers() -> None:
    chart = _chart(summary_header=["Jahr", "Wochen", "Ø Inzidenz"], summary_rows=[["2026", 4, 112.5]])
    sheet = load_workbook(io.BytesIO(build_excel(chart)))["Kennzahlen"]
    rows = [[c.value for c in row][:3] for row in sheet.iter_rows()]
    assert ["2026", 4, 112.5] in rows


def test_pdf_is_a_valid_document_even_with_gaps_in_the_series() -> None:
    pdf = build_pdf(_chart())
    assert pdf.startswith(b"%PDF")
    assert pdf.rstrip().endswith(b"%%EOF")
    assert _page_count(pdf) == 1


def test_pdf_is_still_created_when_the_chart_image_cannot_be_rendered(monkeypatch) -> None:
    def no_browser_engine(chart):
        raise RuntimeError("Kaleido cannot start")

    monkeypatch.setattr(pdf_module, "figure_png", no_browser_engine)
    pdf = build_pdf(_chart())
    assert pdf.startswith(b"%PDF") and pdf.rstrip().endswith(b"%%EOF")


def test_pdf_without_any_value_does_not_fail() -> None:
    assert build_pdf(_chart(series={"A": [None] * 4})).startswith(b"%PDF")


def test_long_tables_continue_on_further_pages() -> None:
    weeks = [f"2025-W{n:02d}" for n in range(1, 53)]
    chart = _chart(categories=weeks, months=["JANUAR"] * 52, series={"A": [float(n) for n in range(52)]})
    assert _page_count(build_pdf(chart)) == 2


def test_figure_shows_the_month_below_the_week_and_thins_out_long_axes() -> None:
    short = build_figure(_chart())
    assert list(short.layout.xaxis.ticktext) == [f"2026-W0{n}<br>JANUAR" for n in range(1, 5)]

    weeks = [f"2024-W{n:02d}" for n in range(1, 53)]
    long = build_figure(_chart(categories=weeks, months=["SEPTEMBER"] * 52, series={"A": [1.0] * 52}))
    ticks = list(long.layout.xaxis.ticktext)
    assert len(ticks) <= 13 and ticks[0] == "2024-W01<br>SEP"  # abbreviated once there are many ticks


def test_figure_draws_reference_series_dashed_and_always_as_lines() -> None:
    figure = build_figure(
        _chart(series={"2026": [1.0] * 4, "Ø 2024, 2025": [2.0] * 4}, reference_series=("Ø 2024, 2025",))
    )
    assert [trace.line.dash for trace in figure.data] == ["solid", "dash"]
    assert {trace.type for trace in figure.data} == {"scatter"}


def test_trend_export_uses_the_dashboard_data(rki_db) -> None:
    data = get_incidence("Baden-Wuerttemberg", "00+", "month")
    chart = trend_export(data, get_status(), NOW)

    assert chart.categories == [p["calendar_week"] for p in data["series"]]
    assert chart.months == ["SEPTEMBER"] * 4
    assert chart.series["ARE-Konsultationsinzidenz"] == [p["value"] for p in data["series"]]
    assert chart.subtitle == "Baden-Württemberg · Alle Altersgruppen · Letzter Monat (4 Wochen) (2026-W36 bis 2026-W39)"
    assert chart.filename == "ARE-Verlauf-Baden-Wuerttemberg-00plus-month-20261001"
    group = chart.kpi_groups[0]
    assert group.title == "Aktuelle Woche · 2026-W39"
    assert [k.title for k in group.items] == ["Aktuell", "Vorwoche", "Veränderung"]
    assert group.items[0].sub == "2026-W39 · je 100.000 Einw."  # every incidence shows its unit
    assert group.items[2].tone in {"up", "down", "flat"}
    assert chart.insight.startswith("Einordnung:")
    assert "Datenstand: 2026-W39" in chart.source_note and "CC BY 4.0" in chart.source_note  # whatever the last ETL was


def test_report_names_the_data_source_of_the_last_import() -> None:
    data = {"region": "Bayern", "region_display_name": "Bayern", "age_group": "00+", "age_group_display_name": "Alle",
            "period": "month", "period_label": "Letzter Monat", "weeks": 0, "from_week": None, "to_week": None,
            "latest_week": None, "latest_value": None, "previous_value": None, "change_percent": None, "series": []}
    live = trend_export(data, {"latest_week": "2026-W39", "last_etl": {"source_type": "RKI_LIVE"}}, NOW)
    test = trend_export(data, {"latest_week": "2026-W39", "last_etl": {"source_type": "LOCAL_TEST"}}, NOW)
    assert "Datenquelle im Monitor: RKI LIVE" in live.source_note
    assert "Datenquelle im Monitor: Testdatei" in test.source_note


def test_year_export_lists_the_years_a_reference_line_and_a_summary(rki_db) -> None:
    data = get_year_comparison("Bundesweit", "00+", [2026, 2025, 2024], "month")
    chart = year_export(data, get_status(), NOW)

    assert chart.categories == ["KW36", "KW37", "KW38", "KW39"]
    assert chart.months == ["SEPTEMBER"] * 4
    assert list(chart.series) == ["2026", "2025", "2024", "Ø 2024, 2025"]  # the reference line names its years
    assert chart.reference_series == ("Ø 2024, 2025",)
    assert all(len(values) == 4 for values in chart.series.values())
    week, window = chart.kpi_groups
    assert week.title == "Aktuelle Woche · 2026-W39"
    assert [k.title for k in week.items] == ["Diese Woche", "Vorjahr", "Veränderung", "Ø Vergleichsjahre"]
    assert week.items[0].sub == "2026-W39 · je 100.000 Einw."  # short title, unit in the small line
    assert week.items[1].sub == "2025-W39 · je 100.000 Einw."
    assert window.title == "Zeitraum · Ø KW36–KW39"
    assert [k.title for k in window.items] == ["2026", "2025", "Veränderung", "Ø Vergleichsjahre", "Abweichung"]
    assert window.items[0].sub == "Ø KW36–KW39 · je 100.000 Einw."
    assert window.items[4].sub == "2026 vom Ø 2024, 2025"
    assert chart.headline and "Vorwoche" in chart.headline
    assert chart.insight.startswith("Einordnung: Die aktuelle Konsultationsinzidenz (2026-W39:")
    assert chart.notes == []
    assert [row[0] for row in chart.summary_rows] == ["2026", "2025", "2024"]
    assert chart.filename == "ARE-Jahresvergleich-Bundesweit-00plus-month-20261001"


@pytest.mark.parametrize("file_format, media_type", [
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("pdf", "application/pdf"),
])
def test_api_exports_trend_and_year_comparison(rki_db, file_format, media_type) -> None:
    trend = api.export_incidence(region="Bayern", age_group="60+", period="quarter", file_format=file_format)
    comparison = api.export_year_comparison(
        region="Bayern", age_group="60+", period="quarter", years=[], file_format=file_format
    )
    for response, prefix in (
        (trend, "ARE-Verlauf-Bayern-60plus-quarter"),
        (comparison, "ARE-Jahresvergleich-Bayern-60plus-quarter"),
    ):
        assert response.media_type == media_type
        assert response.headers["content-disposition"].startswith(f'attachment; filename="{prefix}')
        assert response.headers["content-disposition"].endswith(f'.{file_format}"')
        assert len(response.body) > 1000


def test_api_export_rejects_invalid_requests(rki_db) -> None:
    with pytest.raises(HTTPException) as invalid_period:
        api.export_incidence(region="Bayern", age_group="00+", period="decade", file_format="pdf")
    assert invalid_period.value.status_code == 422

    with pytest.raises(HTTPException) as no_data:
        api.export_incidence(region="Atlantis", age_group="00+", period="month", file_format="pdf")
    assert no_data.value.status_code == 404

    with pytest.raises(HTTPException) as one_year:
        api.export_year_comparison(
            region="Bayern", age_group="00+", period="month", years=[2026], file_format="xlsx"
        )
    assert one_year.value.status_code == 422


def test_state_comparison_export_carries_the_availability_note(rki_db) -> None:
    chart = year_export(get_year_comparison("Bayern", "00+", None, "month"), get_status(), NOW)
    assert any("erst ab 2022-W40" in note for note in chart.notes)
    workbook = load_workbook(io.BytesIO(build_excel(chart)))
    values = [cell.value for row in workbook["Info"].iter_rows() for cell in row if cell.value]
    assert any(isinstance(v, str) and v.startswith("Einordnung") for v in values)
    assert any(isinstance(v, str) and "erst ab 2022-W40" in v for v in values)
    assert build_pdf(chart).startswith(b"%PDF")


def test_api_lists_periods_years_and_comparison(rki_db) -> None:
    assert api.coverage(region="Bayern", age_group="00+").first_week == "2022-W40"
    assert [p.key for p in api.periods()][:2] == ["week", "month"]
    assert api.years(region="Bayern", age_group="00+")[0] == 2026
    result = api.compare_years(region="Bayern", age_group="00+", period="year", years=[])
    assert result.mode == "window" and result.axis[0].label == "KW01"


def test_api_status_reports_the_stored_data(rki_db) -> None:
    status = api.status()
    assert status["total_rows"] > 0
    assert status["latest_week"] == "2026-W39"
    assert status["last_etl"]["status"] == "SUCCESS"
